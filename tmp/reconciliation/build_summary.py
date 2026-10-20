from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
E=ROOT/'tmp/reconciliation/evidence'
OUT=ROOT/'output/docx/Puducherry_Electricity_RTI_Reconciled_Summary.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.63);sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.footer_distance=Inches(.25)
for name,size in [('Normal',10.5),('Title',28),('Subtitle',12),('Heading 1',17),('Heading 2',12),('Caption',9)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7)
 if name=='Normal':st.paragraph_format.line_spacing=1.10
 if name in ['Heading 1','Heading 2']:st.paragraph_format.space_before=Pt(8);st.paragraph_format.keep_with_next=True
 if name=='Title':st.paragraph_format.space_after=Pt(10)
# Suppress theme colors and any inherited paragraph borders.
for st in doc.styles:
 if st.type==1:
  if st.name.startswith(('Heading','Title','Subtitle')):st.font.color.rgb=RGBColor(0,0,0)
  pp=st.element.find(qn('w:pPr'))
  if pp is not None:
   for e in list(pp):
    if e.tag==qn('w:pBdr'):pp.remove(e)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Puducherry electricity RTI summary  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(80,80,80)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Puducherry electricity RTI reconciled summary'
doc.core_properties.subject='Electricity procurement finance infrastructure budgets rooftop solar and smart metering'
doc.core_properties.author='';doc.core_properties.keywords='RTI Puducherry electricity evidence summary'

def p(text='',bold=False,size=None,after=7,keep=False):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(after);x.paragraph_format.keep_with_next=keep
 # ** segments receive emphasis without changing the words.
 for i,s in enumerate(text.split('**')):
  r=x.add_run(s);r.bold=bold or i%2==1
  if size:r.font.size=Pt(size)
 return x

def heading(text):doc.add_heading(text,level=1)
def page(text):doc.add_page_break();heading(text)
def caption(text):
 x=p(text,size=8.5,after=4,keep=True)
 for r in x.runs:r.font.color.rgb=RGBColor(65,65,65)
 return x

def image(name,width=6.7,maxheight=None):
 im=Image.open(E/(name+'.png'));h=width*im.height/im.width
 if maxheight and h>maxheight:width=maxheight*im.width/im.height
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(7)
 shape=x.add_run().add_picture(str(E/(name+'.png')),width=Inches(width))
 shape._inline.docPr.set('descr','Highlighted source extract '+name.replace('_',' '))
 return x

def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.7/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for j,h in enumerate(headers):t.rows[0].cells[j].text=h
 for row in rows:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=str(v)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'233F58' if i==0 else ('EFF4F7' if i%2 else 'FFFFFF'));pr.append(shd)
   borders=OxmlElement('w:tcBorders')
   for k in ['top','left','bottom','right']:
    z=OxmlElement('w:'+k);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'5');z.set(qn('w:color'),'D9D9D9');borders.append(z)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for k,v in [('top',70),('bottom',70),('left',95),('right',95)]:
    z=OxmlElement('w:'+k);z.set(qn('w:w'),str(v));z.set(qn('w:type'),'dxa');m.append(z)
   pr.append(m)
   for x in c.paragraphs:
    x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1.03
    x.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in x.runs:r.font.name='Arial';r.font.size=Pt(9.3);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 p('',after=2)
 return t

# PAGE 1
x=doc.add_paragraph('Puducherry electricity RTI reconciled summary',style='Title')
p('Electricity purchases finances network works solar and smart meters',size=12,after=12)
p('**The records establish purchase cost pressures, greater use of the short term electricity market, unfinished network works and partial solar completion.** They also leave important differences between the electricity and financial figures unexplained.')
caption('Sources A pages 6 and 12  |  B pages 2 and 4 to 10  |  C pages 2 to 4  |  D pages 1 and 2. Highlighted extracts follow each finding.')
heading('The strongest findings')
for text,src in [
 ('Purchase costs exceeded the original approvals by **Rs 860.64 crore across 2022–23 and 2023–24**. The February 2026 reply says FPPCA had not been levied since 2022.','C pages 2 and 3  |  Finding 1'),
 ('Short term market purchases increased from **8.586 MU to 58.893 MU** in 2025–26, almost sevenfold.','A page 12  |  Finding 3'),
 ('The Urban Division spent **Rs 8.07 crore of the Rs 18.79 crore** allocated in its supplied budget tables for 2025–26, about 43 percent.','B pages 8 and 9  |  Finding 5'),
 ('Network strengthening was **37 percent installed** under RDSS; the separate **Rs 472 crore** modernisation proposal was still awaiting sanction.','B pages 4 and 5  |  Finding 4'),
 ('Residential rooftop solar reached **44.3 percent of the 13000 installation target** by 14 September 2026. Government building solar reached **104 of 876 identified buildings**.','B page 10  |  Findings 6 and 7'),
 ('The **Rs 383.59 crore smart meter project** note reports billing integration problems and dates that need correction.','D page 1  |  Findings 9 and 10')]:
 p(text,after=3);caption('Source '+src)
heading('Source documents and dates')
table(['Source','Original PDF','Coverage'],[
 ['A','electricity.pdf','Purchases and payments\nApril 2023 to March 2026'],
 ['B','Electricity_document_1.pdf','Projects and urban accounts for 2025–26\nSolar and status updates to September 2026'],
 ['C','Electricity_document_2.pdf','FPPCA and regulatory costs\nReply dated 6 February 2026'],
 ['D','Smart Meter Project_\nelectricity_document_3.pdf','Smart metering and RDSS loss reduction\nReporting date not established']],widths=[.6,2.7,3.4])
p('**Reading the evidence.** Yellow marks identify the supporting words or figures in the original pages. Page references count PDF pages from the first page. Selected rows or columns are identified in captions. Amounts are rounded only after calculation.',size=9.5)
p('**Units.** 1 MU is 10 lakh electricity units. 1 crore rupees is 100 lakh rupees. MW and kW describe solar system size; they do not measure the units actually generated. Each financial year runs from April to March.',size=9.5)

# PAGE 2
page('1 Purchase costs exceeded the original approvals')
p('The regulator recognised purchase costs above the original tariff allowances for 2022–23 and 2023–24. The combined difference is **Rs 860.64 crore**. Separately, the department states that FPPCA, the adjustment charge for changes in fuel and power purchase costs, had not been levied since 2022, as at its **6 February 2026** reply.')
table(['Year','Original allowance\nRs crore','Recognised actual cost\nRs crore','Difference\nRs crore'],[
 ['2022–23','1307.16','1788.60','481.44'],['2023–24','1480.596','1859.80','379.204']],widths=[1.0,1.85,1.95,1.9])
p('**Meaning for consumers.** The department faced costs above its earlier allowances without collecting that particular adjustment charge. The records do not establish how much of the difference was recovered through other means or remains outstanding. This difference cannot be presented as a household debt.')
caption('Source C page 3 item 5. The original comparison shows the two approved actual cost years. Difference = actual cost minus original allowance.')
image('cost_comparison',width=5.7)
caption('Source C page 2 item 1. Department statement on non levy of FPPCA.')
image('no_fppca',width=6.3)
p('The reply also acknowledges revenue gaps recorded for later recovery. It does not attribute the whole purchase cost difference to non collection of FPPCA.',after=4,size=9.8)
caption('Source C page 4 item 9. A regulatory asset means a revenue gap recorded for recovery later.')
image('revenue_gap',width=6.3)

# PAGE 3
page('2 Electricity quantity figures need reconciliation')
p('The procurement reply reports electricity taken from the central grid. The regulatory reply lists a different annual purchase quantity for the same years. **The difference in coverage is not explained.** The two sets should remain separate until the department supplies a matching statement.')
table(['Year','Central grid drawal\nSource A in MU','Purchase quantity\nSource C in MU','Difference\nin MU'],[
 ['2023–24','3263.63','3597.97','334.34'],['2024–25','3352.72','3649.66','296.94']],widths=[.95,1.95,1.95,1.85])
p('**Meaning.** These are differences between reported quantities. They are not a calculation of electricity losses. Monthly totals in Source A were added to obtain the annual central grid figures above. Source C supplies annual quantities directly.')
caption('Source A page 6 Annexure 3 and Source C page 2 item 3. The yellow columns and entries support the annual comparison.')
image('quantity_comparison',width=6.7,maxheight=4.7)
p('**Date boundary.** Source C was issued before the end of 2025–26. Its 3979 MU entry for that year is therefore excluded from the annual actual comparison above.',size=9.5)
p('Purchase cost recognised by the regulator, payments to suppliers and transmission agencies, and a division’s budget expenditure are also different measures. They cannot be added into one departmental loss figure.',size=9.5)
caption('Financial bases: A page 8 payment table heading; B pages 8 and 9 budget headings; C page 3 item 5 regulatory comparison.')

# PAGE 4
page('3 Short term market purchases increased sharply')
p('Electricity bought through the short term market rose from **8.586 MU in 2024–25 to 58.893 MU in 2025–26**, an increase of about **6.9 times**. The quantity weighted average rate fell slightly, so the rise in estimated cost came from buying much more electricity.')
table(['Year','Quantity\nMU','Weighted rate\nRs per unit','Estimated cost\nRs crore'],[
 ['2023–24','0.837','8.10','0.68'],['2024–25','8.586','6.72','5.77'],['2025–26','58.893','6.40','37.68']],widths=[1.15,1.35,2.0,2.2])
p('**Meaning.** Market buying became more significant. The records do not establish why it increased. The replies contain no matching explanation of the purchase decisions.')
caption('Source A page 12 Annexure 6. Original month quantity and rate columns are reproduced below in three yearly groups; repeated exchange and trader columns are omitted.')
image('market_all',width=6.7)
p('**Calculations.** Annual quantity is the sum of monthly MU. Estimated cost in crore rupees is the sum of MU × monthly rate ÷ 10. Weighted rate is the sum of MU × rate divided by total MU. These use the rounded source entries and are not verified invoice payments.',size=9.5)
p('July 2025 alone records 17.523 MU. July and September together account for approximately 51 percent of 2025–26 market purchases.',size=9.8)
caption('Source A page 12 Annexure 6 July and September 2025 rows highlighted in the final yearly extract above.')

# PAGE 5
page('4 Network works remained partly implemented')
p('The project reply distinguishes approved works from proposals awaiting approval. For RDSS network strengthening, **80 percent of materials had been supplied but only 37 percent of installation was complete**. This means receiving equipment was substantially ahead of installing it. A second RDSS note gives a different snapshot, reconciled in Finding 11.')
caption('Source B page 4 EHV table project 4. Selected original columns retain the project name status funds and physical progress.')
image('rdss',width=6.7)
p('The separate distribution network augmentation and modernisation proposal was estimated at **Rs 472 crore** and was **not sanctioned**. Its proposal had been sent to the Government of India for approval.')
caption('Source B page 5 EHV table project 8. Original name cost sanction and status columns.')
image('modernisation',width=6.6)
p('**Meaning.** The records establish partial implementation and a major proposal awaiting approval. They do not contain the completion dates or performance data needed to measure delay or prove an effect on power cuts or purchase costs.')
p('The EHV table combines “released” and “expended” in one financial column. A release of money to an implementing agency cannot automatically be treated as the value of completed work.',size=9.8)
caption('Source B page 4 financial column heading reproduced in the RDSS extract above.')

# PAGE 6
page('5 Much of the listed Urban Division budget was unspent')
p('The Urban Division’s supplied 2025–26 accounts show **Rs 8.07 crore spent against Rs 18.79 crore allocated**, about **43 percent**. Approximately **Rs 10.72 crore** of that allocation was unspent at March 2026. These are the division’s listed accounts, not a whole department budget.')
table(['Budget type','Allocated\nRs','Spent\nRs','Used'],[
 ['Revenue and running costs','11,89,10,000','5,88,68,942','49.5 percent'],['Capital and infrastructure','6,90,30,000','2,18,78,770','31.7 percent'],['Combined listed entries','18,79,40,000','8,07,47,712','43.0 percent']],widths=[2.6,1.35,1.35,1.4])
p('**Meaning.** Infrastructure spending was particularly low. Repair and maintenance used Rs 46.45 lakh of its Rs 179 lakh allocation, about 26 percent. The accounts do not explain the underspending or measure the resulting condition of the network.',size=10)
caption('Source B page 8 Urban revenue accounts. All 11 monetary rows appear in original order; gaps containing account headings are omitted. Yellow marks show the inputs added for the revenue total.')
image('urban_revenue',width=6.2,maxheight=3.65)
caption('Source B page 9 Urban capital accounts. All five monetary rows appear in original order, with account heading gaps omitted. Figures remain in rupees as printed.')
image('urban_capital',width=6.45,maxheight=2.05)
p('Calculation: sum the allocation and expenditure columns separately. One dash in the revenue expenditure column is treated as no reported spending. Percent used = expenditure ÷ allocation × 100.',size=8.8,after=2)

# PAGE 7
page('6 Residential solar had reached part of its target')
p('By **14 September 2026**, the reply reports **5764 completed residential rooftop installations** against a target of **13000 by March 2027**. That is **44.3 percent**, leaving **7236 installations**. Installed capacity is reported as **20.71 MW**.')
caption('Source B page 10 residential solar opening paragraph. The completion deadline and total installation target are highlighted.')
image('residential_target',width=6.65)
p('The same snapshot records 5526 commissioned installations and 4841 subsidy release cases. There are 238 fewer commissioned installations than completed installations. The records do not identify the processing position or eligibility of the remaining subsidy cases.',size=10)
caption('Source B page 10 residential solar status table. Completion commissioning and subsidy counts use the same 14 September 2026 snapshot.')
image('residential_results',width=6.4,maxheight=3.05)
p('**An internal inconsistency needs clarification.** The paragraph says there is no year wise target, while the table lists a 2025–26 target of 6300 and achievement of 2073, or 32.9 percent. The reply does not explain the difference.',size=10)
caption('Source B page 10 opening statement and annual table. The opening statement is highlighted above; the annual table is reproduced below.')
image('annual_solar',width=6.4)
p('Completed means installation work is recorded as finished. Commissioned is the separate operational milestone used by the reply. These counts should not be treated as interchangeable.',size=9.5)

# PAGE 8
page('7 Government building solar had limited coverage')
p('The government building solar programme was reduced from **1210 proposed buildings to 876 identified buildings**. By **14 September 2026**, **104 buildings** had been solarised, approximately **11.9 percent** of the revised count.')
table(['Measure','Reported position'],[
 ['Original buildings proposed','1210'],['Revised buildings identified','876'],['Buildings solarised','104'],['Proposed capacity','Approximately 27.6 MW'],['Capacity solarised','6.948 MW']],widths=[4.3,2.4])
p('Capacity completed is approximately **25.2 percent** of the proposed 27.6 MW. The building count and capacity percentages differ because buildings have different system sizes. The department attributes the reduction in planned buildings to technical difficulties at sites.')
caption('Source B page 10 government building solar section. Revised scope and recorded completion are highlighted.')
image('government_solar',width=6.7)
p('The same section says agreements signed in 2025–26 covered **105 buildings and 8 MW**. It does not identify whether the 104 completed buildings are exactly the same group. The two counts should not be used to claim that the whole revised programme is nearly complete.')
caption('Source B page 10 agreement paragraph included at the top of the extract above.')
p('**Link to electricity procurement.** Rooftop solar capacity and power bought from solar suppliers describe different sources of electricity. The replies supply no generation or avoided purchase cost statement that would establish the savings from these rooftop systems.')
caption('Related records: B page 10 reports rooftop capacity; A pages 3 to 5 list electricity quantities against generating suppliers. Capacity in MW cannot be added to energy quantities in MU.')

# PAGE 9
page('8 Karaikal solar shows the local implementation position')
p('The Karaikal list covers **26 government buildings with 2015 kW of proposed solar capacity**. Only **one building with 100 kW** is marked completed. Thirteen sites are under progress and twelve are yet to be taken up. This is a local list and should remain separate from the overall programme totals.')
table(['Status','Buildings','Listed capacity'],[
 ['Completed','1','100 kW'],['Under progress','13','1105 kW'],['Yet to be taken up','12','810 kW'],['Total listed','26','2015 kW']],widths=[3.3,1.2,2.2])
p('The completed site is Government Girls Higher Secondary School at Thirunallar. General Hospital Karaikal, with 175 kW proposed, is among the sites yet to be taken up. The count and capacity totals above are calculated from the list.',size=10)
caption('Source B page 2 Karaikal annexure issued with the 22 September 2026 reply. The full site table is reproduced below; the completed status is highlighted.')
image('karaikal',width=6.65,maxheight=5.6)



# SMART METER ADDITION
page('9 Smart metering faced a billing integration obstacle')
p('The new note reports a **Rs 383.59 crore smart metering project** involving PFCCL and Apraava Energy Pvt Ltd. It gives a project award figure rather than an account of payments made. The pages do not supply the total number of meters contracted or an expenditure statement for this component.')
caption('Source D page 1 point 1. The original award amount and contracting parties are shown below. The printed dates are examined in Finding 10.')
image('smart_cost',width=6.7)
p('The note attributes the initial installation approach to opposition to smart meters. It says installations were initially arranged for **government premises, new consumers and solar consumers**, alongside replacement of faulty meters across consumer categories.')
p('The note also states that **NIC, identified in the note as the billing software maintainer, declined support for connecting the smart metering system with the billing software**. Remotely collected readings were instead arranged for regular sharing with the billing system to generate bills. This describes a workaround; the note does not demonstrate full automated billing integration.')
caption('Source D page 1 points 2b and 2c. The initial installation groups integration obstacle and sharing of readings are highlighted.')
image('smart_billing',width=6.7)
p('**Link to the earlier records.** Government and solar consumers overlap with the earlier solar programme topics, but there is no matching list showing which sites received smart meters. The earlier Urban Division meter replacement counts also cannot be treated as smart meter installation totals because they do not identify the meter type.',size=10)
caption('Related records: B page 7 meter replacement entries and B page 10 solar programme figures; D page 1 point 2b above.')
p('The smart metering component and the Rs 106.54 crore RDSS network works are described separately. Their stated costs should not be combined into an expenditure total.',size=9.8)
caption('Source D page 1 point 1 and page 2 final paragraph; the network works figures are reproduced in Finding 11.')

page('10 Meter rollout and feeder monitoring are different milestones')
p('The note reports an installation milestone of **5 percent of smart meters**, followed by testing arrangements and a claimed go live. It supplies no meter count or clear reporting date for that milestone, so **5 percent cannot be presented as the current completion rate**.')
caption('Source D page 1 point 2e. Both the 5 percent milestone and the printed go live date are highlighted.')
image('smart_milestone',width=6.7)
p('**The printed go live date is invalid.** The note says “31.02.2026”; February has no 31st day. It separately prints the contract date as “3.9.2026” and the inauguration date as “1.03.2026”. A corrected and dated project timeline is needed before the sequence or delivery period can be verified.')
caption('Source D page 1 point 1. Contract and inauguration dates as printed. No corrected date has been assumed.')
image('smart_dates',width=6.7)
p('The note separately claims **100 percent feeder integration with the National Feeder Monitoring System**, using an existing electricity monitoring system supplied in 2019, described as SCADA/DMS. Feeders are circuits that carry electricity to groups of consumers. Connecting their monitoring systems is a different task from replacing each consumer’s meter.')
caption('Source D page 1 point 3. The 100 percent figure relates to feeder monitoring integration.')
image('smart_feeders',width=6.55,maxheight=3.05)
p('**Meaning.** The feeder monitoring claim does not establish completion of household smart meter installation or billing integration. Each milestone needs its own dated count and acceptance record.',size=9.8)

page('11 RDSS works and funding need reconciliation')
p('The new note and the earlier project table both identify **Rs 106.54 crore of RDSS network or loss reduction works**. The new note also gives an approved project report estimate of **Rs 83.14 crore** and a central grant sanction of **Rs 49.88 crore**. The award value is **Rs 23.40 crore above the estimate**; the pages do not explain the difference in scope or cost.')
p('For 2025–26 the new note reports **Rs 46.45 crore incurred**, comprising **Rs 13.57 crore from grant funds plus Rs 32.88 crore from the state share**. Those two amounts reconcile exactly. Its 60 percent central funding statement relates to the approved project report cost, rather than the higher works award value.',size=10)
caption('Source D page 2 RDSS loss reduction works. The estimate grant award progress and expenditure figures are highlighted.')
image('rdss_funding',width=6.7,maxheight=3.6)
table(['Measure','Earlier project reply\nSource B','New project note\nSource D'],[
 ['Works value','Rs 106.54 crore','Rs 106.54 crore'],
 ['Materials supplied','80 percent','Around 75 percent'],
 ['Installation completed','37 percent','20 percent'],
 ['2025–26 financial figure','Rs 47 crore\nReleased or expended','Rs 46.45 crore\nIncurred']],widths=[2.3,2.2,2.2])
caption('Source B page 4 project 4. Original works cost funds and physical progress columns are retained for comparison with Source D above.')
image('rdss_comparison',width=6.7,maxheight=1.6)
p('**Meaning.** The matching works value links the two records, but the pages do not establish matching reporting dates and label their financial figures differently. A dated reconciliation is needed to establish progress over time or any unexplained financial difference.',size=9.8)

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
